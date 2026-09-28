"""Generic statement import engine (spec §12, §13, §26, §27, §51, §56, §68).

Workflow: upload → parse (xlsx/csv) → header validation → per-row
transformation + typing + validation + lookups → duplicate detection →
staging rows (`statement_import_rows`) and errors (`statement_import_errors`)
→ user confirms → commit into the *same* normalised tables used by manual
entry and APIs.  Provider differences live only in template rows.
"""
from __future__ import annotations

import csv
import io
import logging
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.core.audit import audit
from app.core.errors import BusinessError, NotFound, OverrideRequired
from app.core.utils import (file_sha256, jsonable, normalize_text, normalize_vehicle_number, now, parse_date,
                            parse_datetime, parse_time, sha256_of, to_decimal, today)
from app.models.finance import BankTransaction
from app.models.imports import (StatementImportBatch, StatementImportError, StatementImportRow,
                                StatementImportTemplate, StatementTemplateColumn, ValueMapping)
from app.models.operations import FuelTransaction, MaintenanceJobCard, MaintenanceType, TollTransaction
from app.models.org import BankAccount, Provider, Vendor
from app.services import documents
from app.services.rules import rule
from app.services.transforms import TContext, apply_pipeline, check_validation, validate_pipeline

log = logging.getLogger("erp.import")

# ───────────────────────── normalised target schemas ─────────────────────────
# field: (data type, required, description)
_BANK = {
    "txn_date": ("DATE", True, "Transaction date"), "value_date": ("DATE", False, "Value date"),
    "narration": ("TEXT", False, "Narration / description"), "credit": ("DECIMAL", False, "Credit amount"),
    "debit": ("DECIMAL", False, "Debit amount"), "amount": ("DECIMAL", False, "Signed amount / amount with type"),
    "dr_cr": ("TEXT", False, "DR/CR indicator (amount mode WITH_TYPE)"),
    "reference_number": ("TEXT", False, "Reference / transaction id"), "utr": ("TEXT", False, "UTR"),
    "cheque_number": ("TEXT", False, "Cheque number"), "balance": ("DECIMAL", False, "Running balance"),
    "counterparty": ("TEXT", False, "Counterparty name"),
}
_FUEL = {
    "provider_transaction_id": ("TEXT", False, "Provider transaction id"),
    "txn_datetime": ("DATETIME", False, "Date & time"), "txn_date": ("DATE", False, "Date"),
    "txn_time": ("TIME", False, "Time"), "vehicle_number": ("TEXT", False, "Vehicle number"),
    "card_number": ("TEXT", False, "Fuel / fleet card"), "station_code": ("TEXT", False, "Station code"),
    "station_name": ("TEXT", False, "Station name"), "fuel_type": ("TEXT", True, "Fuel type / product"),
    "quantity": ("DECIMAL", True, "Quantity"), "unit": ("TEXT", False, "Unit"), "rate": ("DECIMAL", False, "Rate"),
    "amount": ("DECIMAL", False, "Amount"), "tax_amount": ("DECIMAL", False, "Tax"),
    "discount_amount": ("DECIMAL", False, "Discount"), "total_amount": ("DECIMAL", False, "Total"),
    "invoice_number": ("TEXT", False, "Invoice"), "odometer": ("DECIMAL", False, "Odometer"),
    "remarks": ("TEXT", False, "Remarks"),
}
_TOLL = {
    "transaction_id": ("TEXT", True, "Provider transaction id"), "txn_datetime": ("DATETIME", False, "Date & time"),
    "txn_date": ("DATE", False, "Date"), "txn_time": ("TIME", False, "Time"),
    "vehicle_number": ("TEXT", False, "Vehicle number"), "fastag_id": ("TEXT", False, "FASTag id"),
    "plaza_id": ("TEXT", False, "Toll plaza id"), "plaza_code": ("TEXT", False, "Toll plaza code"),
    "plaza_name": ("TEXT", False, "Toll plaza / gate name"), "amount": ("DECIMAL", True, "Amount"),
    "direction": ("TEXT", False, "Direction"), "lane": ("TEXT", False, "Lane"),
    "txn_kind": ("TEXT", False, "DEBIT / REFUND / RECHARGE"), "remarks": ("TEXT", False, "Remarks"),
}
_MAINT = {
    "job_card_number": ("TEXT", True, "Job card / bill number"), "job_date": ("DATE", True, "Date"),
    "vehicle_number": ("TEXT", True, "Vehicle number"), "maintenance_type": ("TEXT", True, "Maintenance type"),
    "odometer": ("DECIMAL", False, "Odometer"), "complaint": ("TEXT", False, "Complaint"),
    "work_performed": ("TEXT", False, "Work performed"), "vendor_name": ("TEXT", False, "Vendor"),
    "parts_amount": ("DECIMAL", False, "Parts"), "labour_amount": ("DECIMAL", False, "Labour"),
    "other_amount": ("DECIMAL", False, "Other"), "tax_amount": ("DECIMAL", False, "Tax"),
    "invoice_number": ("TEXT", False, "Invoice"),
}
_INS = {
    "policy_number": ("TEXT", True, "Policy number"), "vehicle_number": ("TEXT", False, "Vehicle number"),
    "policy_type": ("TEXT", True, "Policy type"), "start_date": ("DATE", True, "Start"),
    "expiry_date": ("DATE", True, "Expiry"), "insured_amount": ("DECIMAL", False, "Sum insured / IDV"),
    "premium": ("DECIMAL", False, "Premium"), "tax_amount": ("DECIMAL", False, "Tax"),
}
_GPS = {"vehicle_number": ("TEXT", True, "Vehicle number"), "reading_datetime": ("DATETIME", True, "Reading time"),
        "odometer": ("DECIMAL", True, "Odometer")}
TARGET_FIELDS: dict[str, dict] = {"BANK": _BANK, "FUEL": _FUEL, "FUEL_CARD": _FUEL, "TOLL": _TOLL, "FASTAG": _TOLL,
                                  "MAINTENANCE": _MAINT, "INSURANCE": _INS, "GPS": _GPS, "TELEMATICS": _GPS}
LOOKUPS = {"vehicle", "fuel_type", "fuel_station", "toll_plaza", "maintenance_type", "vendor", "provider"}
FAMILY = {"FUEL_CARD": "FUEL", "FASTAG": "TOLL", "TELEMATICS": "GPS"}


def family(statement_type: str) -> str:
    return FAMILY.get(statement_type, statement_type)


# ───────────────────────── template validation (used by the template builder) ─────────────────────────
def template_after_save(ctx, tpl: StatementImportTemplate, data: dict, is_new: bool) -> None:
    fields = TARGET_FIELDS.get(tpl.statement_type)
    if fields is None:
        raise BusinessError(f"Unsupported statement type {tpl.statement_type}")
    if tpl.data_start_row <= tpl.header_row:
        raise BusinessError("Data start row must be after the header row")
    cols = ctx.db.execute(select(StatementTemplateColumn).where(StatementTemplateColumn.template_id == tpl.id)
                          ).scalars().all()
    errs = []
    for c in cols:
        if c.target_field not in fields:
            errs.append(f"Unknown target field '{c.target_field}' for {tpl.statement_type}. "
                        f"Allowed: {', '.join(fields)}")
        if not c.source_column and not c.source_header and not c.default_value:
            errs.append(f"{c.target_field}: give a source column letter, a source header or a default value")
        if c.source_column and not c.source_column.strip().isdigit():
            try:
                column_index_from_string(c.source_column.strip().upper())
            except ValueError:
                errs.append(f"{c.target_field}: invalid column '{c.source_column}'")
        try:
            validate_pipeline(c.transformation)
        except ValueError as exc:
            errs.append(f"{c.target_field}: {exc}")
        if c.lookup_rule and c.lookup_rule not in LOOKUPS:
            errs.append(f"{c.target_field}: unknown lookup '{c.lookup_rule}' (allowed: {', '.join(sorted(LOOKUPS))})")
    mapped = {c.target_field for c in cols}
    fam = family(tpl.statement_type)
    if fam == "BANK":
        need = {"SEPARATE": {"credit", "debit"}, "SIGNED": {"amount"}, "WITH_TYPE": {"amount", "dr_cr"}}[tpl.amount_mode]
        if not need <= mapped:
            errs.append(f"Amount mode {tpl.amount_mode} requires columns: {', '.join(sorted(need))}")
    if fam in ("FUEL", "TOLL") and not ({"txn_datetime", "txn_date"} & mapped):
        errs.append("Map txn_datetime or txn_date")
    for f, (_, req, _) in fields.items():
        if req and f not in mapped:
            errs.append(f"Required target field '{f}' is not mapped")
    if errs:
        raise BusinessError("Template configuration errors", "TEMPLATE",
                            [{"field": "columns", "message": e} for e in errs], 422)


# ───────────────────────── parsing ─────────────────────────
@dataclass
class Parsed:
    sheet: str | None
    sheets: list[str]
    header: list[Any]
    rows: list[tuple[int, list[Any]]] = field(default_factory=list)


def list_sheets(data: bytes, filename: str) -> list[str]:
    if filename.lower().endswith(".csv"):
        return []
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def parse_file(data: bytes, filename: str, tpl: StatementImportTemplate, sheet: str | None = None) -> Parsed:
    name = filename.lower()
    if name.endswith(".xls"):
        raise BusinessError("Legacy .xls files are not supported; save the file as .xlsx or .csv")
    blank_limit = max(tpl.stop_at_blank_rows or 3, 1)
    footer = [normalize_text(k) for k in (tpl.skip_footer_keywords or "").split(",") if k.strip()]
    rows: list[tuple[int, list[Any]]] = []
    header: list[Any] = []
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        it = enumerate(csv.reader(io.StringIO(text)), start=1)
        sheets, used = [], None
    else:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets = wb.sheetnames
        used = sheet or tpl.sheet_name or sheets[0]
        if used not in sheets:
            wb.close()
            raise BusinessError(f"Sheet '{used}' not found. Available: {', '.join(sheets)}")
        it = enumerate(wb[used].iter_rows(values_only=True), start=1)
    blanks = 0
    for rn, values in it:
        values = list(values)
        if rn == tpl.header_row:
            header = values
        if rn < tpl.data_start_row:
            continue
        if all(v is None or str(v).strip() == "" for v in values):
            blanks += 1
            if blanks >= blank_limit:
                break
            continue
        blanks = 0
        first = next((normalize_text(v) for v in values if v not in (None, "")), "")
        if footer and any(first.startswith(k) for k in footer):
            continue
        rows.append((rn, values))
    if not name.endswith(".csv"):
        wb.close()
    return Parsed(used, sheets, header, rows)


def _col_index(c: StatementTemplateColumn, header_map: dict[str, int]) -> int | None:
    if c.source_column:
        s = c.source_column.strip().upper()
        return int(s) - 1 if s.isdigit() else column_index_from_string(s) - 1
    if c.source_header:
        return header_map.get(normalize_text(c.source_header))
    return None


def validate_headers(tpl: StatementImportTemplate, cols: list[StatementTemplateColumn], header: list[Any]
                     ) -> tuple[list[dict], dict[int, int | None]]:
    header_map = {normalize_text(h): i for i, h in enumerate(header) if h not in (None, "")}
    problems, idx = [], {}
    for c in cols:
        i = _col_index(c, header_map)
        idx[c.id] = i
        if c.source_header:
            if i is None:
                problems.append({"column": c.source_column or "", "field": c.target_field, "severity":
                                 "ERROR" if c.is_required else "WARNING",
                                 "message": f"Header '{c.source_header}' not found in row {tpl.header_row}"})
            elif c.source_column:
                actual = header[i] if i < len(header) else None
                if normalize_text(actual) != normalize_text(c.source_header):
                    problems.append({"column": c.source_column, "field": c.target_field, "severity": "ERROR",
                                     "message": f"Column {c.source_column} header is '{actual}', expected "
                                                f"'{c.source_header}'. Is this the right template/version?",
                                     "original": actual})
    return problems, idx


# ───────────────────────── row normalisation ─────────────────────────
class RowResult:
    def __init__(self, rn: int, raw: dict):
        self.rn = rn
        self.raw = raw
        self.values: dict[str, Any] = {}
        self.issues: list[dict] = []
        self.status = "VALID"
        self.key: str | None = None

    def issue(self, severity, etype, msg, field_=None, column=None, original=None, suggestion=None):
        self.issues.append({"severity": severity, "error_type": etype, "message": msg, "field": field_,
                            "column": column, "original": None if original is None else str(original)[:500],
                            "suggestion": suggestion})
        if severity == "ERROR":
            self.status = "ERROR"
        elif severity == "WARNING" and self.status == "VALID":
            self.status = "WARNING"


def _typed(dtype: str, v: Any, tpl: StatementImportTemplate) -> Any:
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    if dtype == "DATE":
        return v if isinstance(v, date) and not isinstance(v, datetime) else parse_date(v, tpl.date_format)
    if dtype == "DATETIME":
        if isinstance(v, datetime):
            return v
        d = parse_datetime(v, tpl.datetime_format)
        if d is None:
            raise ValueError(f"Invalid date/time '{v}'")
        return d
    if dtype == "TIME":
        return parse_time(v)
    if dtype == "DECIMAL":
        return v if isinstance(v, Decimal) else to_decimal(v)
    if dtype == "INTEGER":
        d = to_decimal(v)
        return int(d) if d is not None else None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


class Lookups:
    """Cached lookups for one batch (vehicle, fuel type, station, plaza ...)."""

    def __init__(self, db: Session, provider_id: int):
        self.db, self.pid, self.cache = db, provider_id, {}

    def mapper(self, mtype: str, value: str) -> str | None:
        k = ("map", mtype, normalize_text(value))
        if k not in self.cache:
            from app.services.operations import mapped_value
            vm = mapped_value(self.db, mtype, self.pid, value)
            self.cache[k] = (vm.target_value or (str(vm.target_id) if vm.target_id else None)) if vm else None
        return self.cache[k]

    def resolve(self, kind: str, value: Any, row: dict) -> int | None:
        k = (kind, normalize_text(value))
        if k in self.cache:
            return self.cache[k]
        from app.services import operations as ops
        db, pid = self.db, self.pid
        rid = None
        if kind == "vehicle":
            rid = ops.match_vehicle(db, pid, str(value), row.get("txn_date") or row.get("job_date"),
                                    fastag=row.get("fastag_id"), card=row.get("card_number"))
        elif kind == "fuel_type":
            rid = ops.match_fuel_type(db, pid, str(value))
        elif kind == "fuel_station":
            rid = ops.match_station(db, pid, row.get("station_code"), row.get("station_name") or str(value))
        elif kind == "toll_plaza":
            rid = ops.match_plaza(db, pid, row.get("plaza_id"), row.get("plaza_code"), row.get("plaza_name"))
        elif kind == "maintenance_type":
            s = normalize_text(value)
            r = db.execute(select(MaintenanceType.id).where(
                (func.upper(MaintenanceType.code) == s) | (func.upper(MaintenanceType.name) == s))).first()
            rid = r[0] if r else None
        elif kind == "vendor":
            r = db.execute(select(Vendor.id).where(func.upper(Vendor.name) == normalize_text(value))).first()
            rid = r[0] if r else None
        elif kind == "provider":
            r = db.execute(select(Provider.id).where((func.upper(Provider.name) == normalize_text(value)) |
                                                     (func.upper(Provider.code) == normalize_text(value)))).first()
            rid = r[0] if r else None
        if kind in ("vehicle",) and rid is None:
            pass  # never guessed: remains unmatched for review
        self.cache[k] = rid
        return rid


def normalize_row(rn: int, values: list[Any], tpl: StatementImportTemplate, cols: list[StatementTemplateColumn],
                  idx: dict[int, int | None], lk: Lookups) -> RowResult:
    raw = {get_column_letter(i + 1): v for i, v in enumerate(values)}
    r = RowResult(rn, raw)
    fields = TARGET_FIELDS[tpl.statement_type]
    ctxrow = dict(raw)
    ctx = TContext(ctxrow, lk.mapper, tpl.date_format, tpl.datetime_format)
    for c in sorted(cols, key=lambda c: c.display_order):
        i = idx.get(c.id)
        letter = get_column_letter(i + 1) if i is not None else (c.source_column or "")
        original = values[i] if i is not None and i < len(values) else None
        try:
            v = apply_pipeline(c.transformation, original, ctx)
            if (v is None or v == "") and c.default_value not in (None, ""):
                v = c.default_value
            schema_type = fields.get(c.target_field, ("TEXT",))[0]
            # the normalised schema's type wins unless the column explicitly asks for another non-text type
            dtype = c.data_type if c.data_type and c.data_type != "TEXT" else schema_type
            v = _typed(dtype, v, tpl)
        except (ValueError, ArithmeticError) as exc:
            r.issue("ERROR", "FORMAT", str(exc), c.target_field, letter, original,
                    "Use DD/MM/YYYY for dates; numbers without text" if "date" in str(exc).lower() else None)
            continue
        if v in (None, "") and c.is_required:
            r.issue("ERROR", "REQUIRED", f"{c.target_field} is required", c.target_field, letter, original)
            continue
        msg = check_validation(c.validation_rule, v)
        if msg:
            r.issue("ERROR", "VALIDATION", f"{c.target_field} {msg}", c.target_field, letter, original)
            continue
        if v not in (None, "") or c.target_field not in r.values:
            r.values[c.target_field] = v
        r.values.setdefault("_cols", {})[c.target_field] = letter
    for c in cols:
        if c.lookup_rule and r.values.get(c.target_field) not in (None, ""):
            rid = lk.resolve(c.lookup_rule, r.values[c.target_field], r.values)
            r.values[f"{c.target_field}__id"] = rid
            if rid is None:
                sev = "ERROR" if family(tpl.statement_type) in ("MAINTENANCE", "GPS") and c.lookup_rule in (
                    "vehicle", "maintenance_type") else "WARNING"
                r.issue(sev, "UNMATCHED", f"No {c.lookup_rule.replace('_', ' ')} matches '{r.values[c.target_field]}'",
                        c.target_field, r.values["_cols"].get(c.target_field), r.values[c.target_field],
                        "Add the master record or a value mapping, then re-match")
    _finalize(tpl, r)
    return r


def _finalize(tpl: StatementImportTemplate, r: RowResult) -> None:
    v = r.values
    fam = family(tpl.statement_type)
    if r.status == "ERROR":
        return
    if fam == "BANK":
        cr, dr = v.get("credit") or Decimal(0), v.get("debit") or Decimal(0)
        if tpl.amount_mode == "SIGNED":
            a = v.get("amount") or Decimal(0)
            cr, dr = (a, Decimal(0)) if a > 0 else (Decimal(0), -a)
        elif tpl.amount_mode == "WITH_TYPE":
            a = abs(v.get("amount") or Decimal(0))
            ind = normalize_text(v.get("dr_cr"))
            if ind in ("CR", "C", "CREDIT", "DEPOSIT"):
                cr, dr = a, Decimal(0)
            elif ind in ("DR", "D", "DEBIT", "WITHDRAWAL"):
                cr, dr = Decimal(0), a
            else:
                r.issue("ERROR", "VALIDATION", f"Unknown DR/CR indicator '{v.get('dr_cr')}'", "dr_cr",
                        v.get("_cols", {}).get("dr_cr"), v.get("dr_cr"))
                return
        cr, dr = abs(cr), abs(dr)
        if cr and dr:
            r.issue("ERROR", "VALIDATION", "Row has both debit and credit", "credit")
        elif not cr and not dr:
            r.issue("WARNING", "VALIDATION", "Row has no amount (e.g. opening/closing balance line) — skipped", "amount")
            r.status = "SKIPPED"
        v["credit"], v["debit"] = cr, dr
    elif fam in ("FUEL", "TOLL"):
        dt = v.get("txn_datetime")
        if dt is None and v.get("txn_date"):
            dt = datetime.combine(v["txn_date"], v.get("txn_time") or time(0))
        if dt is None:
            r.issue("ERROR", "REQUIRED", "Transaction date is required", "txn_date")
            return
        v["txn_datetime"], v["txn_date"] = dt, dt.date()
        if fam == "FUEL":
            if not v.get("quantity") or v["quantity"] <= 0:
                r.issue("ERROR", "VALIDATION", "Quantity must be greater than zero", "quantity")
            if v.get("amount") is None and v.get("total_amount") is None and v.get("rate") is None:
                r.issue("ERROR", "REQUIRED", "Amount, total or rate is required", "amount")
            if v.get("amount") is None and v.get("total_amount") is not None:
                v["amount"] = v["total_amount"] - (v.get("tax_amount") or 0) + (v.get("discount_amount") or 0)
        else:
            kind = normalize_text(v.get("txn_kind") or "DEBIT")
            v["txn_kind"] = "REFUND" if kind in ("REFUND", "CREDIT", "REVERSAL", "CR") else \
                "RECHARGE" if kind in ("RECHARGE", "TOPUP", "TOP UP") else "DEBIT"
            if v.get("amount") is not None and v["amount"] < 0:
                v["amount"] = -v["amount"]
                if v["txn_kind"] == "DEBIT":
                    v["txn_kind"] = "REFUND"
    elif fam == "GPS":
        pass


def _business_key(tpl: StatementImportTemplate, batch: StatementImportBatch, v: dict, occurrence: int) -> str:
    fam = family(tpl.statement_type)
    if tpl.duplicate_key_fields:
        parts = [v.get(f.strip()) for f in tpl.duplicate_key_fields.split(",")]
        return sha256_of(fam, batch.provider_id, *parts, occurrence if fam == "BANK" else 0)
    if fam == "BANK":
        from app.services.bank import bank_hash
        return bank_hash(batch.bank_account_id, v["txn_date"], v.get("narration"), v["debit"], v["credit"],
                         v.get("reference_number"), v.get("utr"), v.get("balance"), occurrence)
    if fam == "TOLL":
        return sha256_of("TOLL", batch.provider_id, v.get("transaction_id"))
    if fam == "FUEL":
        if v.get("provider_transaction_id"):
            return sha256_of("FUELID", batch.provider_id, v["provider_transaction_id"])
        return sha256_of("FUEL", batch.provider_id, None, v.get("vehicle_number__id") or normalize_vehicle_number(
            v.get("vehicle_number")), v.get("txn_datetime"), v.get("quantity"), v.get("amount"))
    if fam == "MAINTENANCE":
        return sha256_of("MAINT", v.get("job_card_number"))
    if fam == "INSURANCE":
        return sha256_of("INS", batch.provider_id, v.get("policy_number"))
    return sha256_of("GPS", v.get("vehicle_number__id"), v.get("reading_datetime"))


def _existing_keys(db: Session, tpl, batch, keys: list[str], rows: list[RowResult]) -> set[str]:
    fam = family(tpl.statement_type)
    found: set[str] = set()
    chunk = 500
    if fam == "BANK":
        for i in range(0, len(keys), chunk):
            found |= set(db.execute(select(BankTransaction.txn_hash).where(
                BankTransaction.bank_account_id == batch.bank_account_id,
                BankTransaction.txn_hash.in_(keys[i:i + chunk]))).scalars())
    elif fam == "TOLL":
        ids = {r.values.get("transaction_id"): r.key for r in rows if r.key}
        tid = list(ids)
        for i in range(0, len(tid), chunk):
            for t in db.execute(select(TollTransaction.transaction_id).where(
                    TollTransaction.provider_id == batch.provider_id,
                    TollTransaction.transaction_id.in_(tid[i:i + chunk]))).scalars():
                found.add(ids[t])
    elif fam == "FUEL":
        with_id = {r.values.get("provider_transaction_id"): r.key for r in rows if r.values.get("provider_transaction_id")}
        tid = list(with_id)
        for i in range(0, len(tid), chunk):
            for t in db.execute(select(FuelTransaction.provider_transaction_id).where(
                    FuelTransaction.provider_id == batch.provider_id,
                    FuelTransaction.provider_transaction_id.in_(tid[i:i + chunk]))).scalars():
                found.add(with_id[t])
        others = [r.key for r in rows if r.key and not r.values.get("provider_transaction_id")]
        from app.services.operations import fuel_hash  # noqa: F401 - same hash family
        for i in range(0, len(others), chunk):
            found |= set(db.execute(select(FuelTransaction.txn_hash).where(
                FuelTransaction.txn_hash.in_(others[i:i + chunk]))).scalars())
    elif fam == "MAINTENANCE":
        nums = {r.values.get("job_card_number"): r.key for r in rows if r.key}
        for n in db.execute(select(MaintenanceJobCard.job_card_number).where(
                MaintenanceJobCard.job_card_number.in_(list(nums)))).scalars():
            found.add(nums[n])
    elif fam == "INSURANCE":
        from app.models.compliance import InsurancePolicy
        nums = {r.values.get("policy_number"): r.key for r in rows if r.key}
        for n in db.execute(select(InsurancePolicy.policy_number).where(
                InsurancePolicy.provider_id == batch.provider_id,
                InsurancePolicy.policy_number.in_(list(nums)))).scalars():
            found.add(nums[n])
    return found


# ───────────────────────── preview ─────────────────────────
def create_preview(ctx, template_id: int, data: bytes, filename: str, sheet: str | None = None,
                   bank_account_id: int | None = None) -> StatementImportBatch:
    db = ctx.db
    tpl = db.get(StatementImportTemplate, template_id)
    if not tpl or not tpl.is_active:
        raise NotFound("Active import template")
    if tpl.effective_to and tpl.effective_to < today():
        raise BusinessError("This template version has expired; select the current version")
    fam = family(tpl.statement_type)
    if fam == "BANK":
        acct = db.get(BankAccount, bank_account_id) if bank_account_id else None
        if not acct or not acct.is_active:
            raise BusinessError("Select the bank account this statement belongs to")
    fname = documents.safe_filename(filename)
    if not fname.lower().endswith((".xlsx", ".csv", ".xls")):
        raise BusinessError("Upload an Excel (.xlsx) or CSV file")
    documents.validate_upload(fname, data) if not fname.lower().endswith(".csv") else None
    fhash = file_sha256(data)
    prev = db.execute(select(StatementImportBatch).where(
        StatementImportBatch.file_hash == fhash, StatementImportBatch.template_id == tpl.id,
        StatementImportBatch.status.in_(["COMPLETED", "COMPLETED_WITH_ERRORS"]))).scalars().first()
    if prev and rule(db, "DUPLICATE_FILE_CHECK") == "BLOCK":
        raise BusinessError(f"This exact file was already imported in batch #{prev.id} on "
                            f"{prev.imported_at:%d/%m/%Y}", "DUPLICATE_FILE")
    cols = list(db.execute(select(StatementTemplateColumn).where(StatementTemplateColumn.template_id == tpl.id)
                           .order_by(StatementTemplateColumn.display_order)).scalars())
    if not cols:
        raise BusinessError("The template has no column mappings")
    parsed = parse_file(data, fname, tpl, sheet)
    batch = StatementImportBatch(provider_id=tpl.provider_id, template_id=tpl.id, statement_type=tpl.statement_type,
                                 bank_account_id=bank_account_id if fam == "BANK" else None, file_name=fname,
                                 file_hash=fhash, storage_key=documents.store_bytes(data, "." + fname.rsplit(".", 1)[-1]),
                                 sheet_name=parsed.sheet, status="PREVIEWED", total_rows=len(parsed.rows),
                                 created_by=ctx.user.id, updated_by=ctx.user.id,
                                 summary={"sheets": parsed.sheets, "duplicate_file_of": prev.id if prev else None})
    db.add(batch)
    db.flush()
    header_problems, idx = validate_headers(tpl, cols, parsed.header)
    for hp in header_problems:
        db.add(StatementImportError(batch_id=batch.id, source_row=tpl.header_row, source_column=hp["column"],
                                    target_field=hp["field"], original_value=hp.get("original"), error_type="HEADER",
                                    error_message=hp["message"], severity=hp["severity"]))
    header_ok = not any(h["severity"] == "ERROR" for h in header_problems)
    if prev:
        db.add(StatementImportError(batch_id=batch.id, error_type="DUPLICATE", severity="WARNING",
                                    error_message=f"Same file was already imported in batch #{prev.id}. "
                                                  "Rows already present will be flagged as duplicates."))
    lk = Lookups(db, tpl.provider_id)
    results: list[RowResult] = []
    if header_ok:
        occ: Counter = Counter()
        for rn, values in parsed.rows:
            r = normalize_row(rn, values, tpl, cols, idx, lk)
            if r.status not in ("ERROR", "SKIPPED"):
                base = _business_key(tpl, batch, r.values, 0)
                occurrence = occ[base]
                occ[base] += 1
                r.key = _business_key(tpl, batch, r.values, occurrence) if fam == "BANK" else base
                if fam != "BANK" and occurrence > 0:
                    r.issue("WARNING", "DUPLICATE", "Duplicate of an earlier row in this file", None)
                    r.status = "DUPLICATE"
            results.append(r)
        existing = _existing_keys(db, tpl, batch, [r.key for r in results if r.key], results)
        for r in results:
            if r.key and r.key in existing and r.status != "DUPLICATE":
                r.issue("WARNING", "DUPLICATE", "Already imported (same business key) — will not be inserted again")
                r.status = "DUPLICATE"
    _store_staging(db, batch, results)
    _count(db, batch)
    batch.summary = {**(batch.summary or {}), "header_ok": header_ok, "columns": [
        {"target": c.target_field, "source": c.source_column or c.source_header, "transformation": c.transformation,
         "lookup": c.lookup_rule} for c in cols]}
    audit(db, ctx.user, "IMPORT_PREVIEW", "statement_import_batches", batch.id,
          new={"file": fname, "rows": batch.total_rows, "template": tpl.template_name})
    return batch


def _store_staging(db: Session, batch: StatementImportBatch, results: list[RowResult]) -> None:
    objs, errs = [], []
    for r in results:
        vals = {k: v for k, v in r.values.items() if k != "_cols"}
        objs.append(StatementImportRow(batch_id=batch.id, source_row=r.rn, raw_values=jsonable(r.raw),
                                       normalized_values=jsonable(vals), row_hash=r.key, status=r.status,
                                       messages="; ".join(i["message"] for i in r.issues) or None))
        for i in r.issues:
            errs.append(StatementImportError(batch_id=batch.id, source_row=r.rn, source_column=i["column"],
                                             target_field=i["field"], original_value=i["original"],
                                             error_type=i["error_type"], error_message=i["message"][:500],
                                             suggested_correction=i["suggestion"], severity=i["severity"]))
    for i in range(0, len(objs), 1000):
        db.add_all(objs[i:i + 1000])
        db.flush()
    db.add_all(errs)
    db.flush()


def _count(db: Session, batch: StatementImportBatch) -> None:
    counts = dict(db.execute(select(StatementImportRow.status, func.count()).where(
        StatementImportRow.batch_id == batch.id).group_by(StatementImportRow.status)).all())
    batch.error_rows = counts.get("ERROR", 0)
    batch.duplicate_rows = counts.get("DUPLICATE", 0)
    batch.warning_rows = counts.get("WARNING", 0)
    batch.successful_rows = counts.get("IMPORTED", 0)
    batch.total_rows = sum(counts.values())


# ───────────────────────── commit ─────────────────────────
def _decode(v: dict) -> dict:
    """Staging JSON → Python types for committing."""
    out = {}
    for k, x in v.items():
        if isinstance(x, str) and k in ("txn_date", "value_date", "job_date", "start_date", "expiry_date"):
            out[k] = date.fromisoformat(x)
        elif isinstance(x, str) and k in ("txn_datetime", "reading_datetime"):
            out[k] = datetime.fromisoformat(x)
        elif isinstance(x, str) and k == "txn_time":
            out[k] = time.fromisoformat(x)
        elif isinstance(x, str) and k in ("credit", "debit", "amount", "balance", "quantity", "rate", "tax_amount",
                                          "discount_amount", "total_amount", "odometer", "parts_amount",
                                          "labour_amount", "other_amount", "insured_amount", "premium"):
            out[k] = Decimal(x)
        else:
            out[k] = x
    return out


def _src(batch: StatementImportBatch, row: StatementImportRow) -> dict:
    return {"source_type": "EXCEL", "source_file": batch.file_name, "source_sheet": batch.sheet_name,
            "source_row": row.source_row, "import_batch_id": batch.id,
            "source_reference": f"BATCH#{batch.id}/ROW{row.source_row}"}


def commit_batch(ctx, batch_id: int, auto_suggest: bool = True, progress_cb=None) -> StatementImportBatch:
    db = ctx.db
    batch = db.get(StatementImportBatch, batch_id, with_for_update=True)
    if not batch:
        raise NotFound("Import batch")
    if batch.status not in ("PREVIEWED", "QUEUED"):
        raise BusinessError(f"Batch is {batch.status}; only previewed batches can be committed")
    if not (batch.summary or {}).get("header_ok", True):
        raise BusinessError("Header validation failed; fix the template or the file and upload again")
    tpl = db.get(StatementImportTemplate, batch.template_id)
    fam = family(tpl.statement_type)
    batch.status = "PROCESSING"
    rows = db.execute(select(StatementImportRow).where(StatementImportRow.batch_id == batch.id,
                                                       StatementImportRow.status.in_(["VALID", "WARNING"]))
                      .order_by(StatementImportRow.source_row)).scalars().all()
    created_bank: list[int] = []
    unmatched = 0
    for n, row in enumerate(rows, 1):
        v = _decode(row.normalized_values or {})
        try:
            with db.begin_nested():
                obj = _commit_row(ctx, fam, batch, tpl, row, v)
        except OverrideRequired as exc:
            row.status, row.messages = "ERROR", exc.message
            db.add(StatementImportError(batch_id=batch.id, source_row=row.source_row, error_type="VALIDATION",
                                        error_message=exc.message[:500], severity="ERROR"))
            continue
        except BusinessError as exc:
            row.status, row.messages = "ERROR", exc.message
            db.add(StatementImportError(batch_id=batch.id, source_row=row.source_row, error_type="VALIDATION",
                                        error_message=exc.message[:500], severity="ERROR"))
            continue
        row.status, row.target_table, row.target_id = "IMPORTED", obj.__tablename__, obj.id
        if fam == "BANK":
            created_bank.append(obj.id)
        if getattr(obj, "status", "") in ("VEHICLE_UNMATCHED", "PLAZA_UNMATCHED", "VEHICLE_MATCHED", "FLAGGED"):
            unmatched += 1
        if progress_cb and n % 200 == 0:
            progress_cb(n, len(rows))
    batch.unmatched_rows = unmatched
    batch.processed_rows = len(rows)
    db.flush()
    _count(db, batch)
    batch.imported_at, batch.imported_by = now(), ctx.user.id
    batch.status = "COMPLETED" if not batch.error_rows else "COMPLETED_WITH_ERRORS"
    if fam == "BANK":
        batch.unmatched_rows = len(created_bank)
        if created_bank and auto_suggest:
            from app.services.bank import auto_suggest as _auto
            res = _auto(ctx, created_bank)
            batch.summary = {**(batch.summary or {}), "auto_suggest": res}
        acct = db.get(BankAccount, batch.bank_account_id)
        last = db.execute(select(BankTransaction).where(BankTransaction.import_batch_id == batch.id,
                                                        BankTransaction.balance.is_not(None))
                          .order_by(BankTransaction.txn_date.desc(), BankTransaction.source_row.desc())).scalars().first()
        if last and (acct.closing_balance_date is None or last.txn_date >= acct.closing_balance_date):
            acct.closing_balance, acct.closing_balance_date = last.balance, last.txn_date
    audit(db, ctx.user, "IMPORT", "statement_import_batches", batch.id, new={
        "imported": batch.successful_rows, "duplicates": batch.duplicate_rows, "errors": batch.error_rows})
    from app.services.notifications import notify
    if batch.error_rows:
        notify(db, "IMPORT_ERRORS", f"Import #{batch.id} ({batch.file_name}) has {batch.error_rows} error row(s)",
               severity="WARNING", entity_type="statement_import_batches", entity_id=batch.id,
               link_url=f"/imports/batches/{batch.id}", dedupe_key=f"import-errors:{batch.id}")
    return batch


def _commit_row(ctx, fam: str, batch, tpl, row, v: dict):
    db = ctx.db
    src = _src(batch, row)
    if fam == "BANK":
        from app.services.bank import _hist, recompute
        t = BankTransaction(bank_account_id=batch.bank_account_id, txn_date=v["txn_date"], value_date=v.get("value_date"),
                            narration=(v.get("narration") or "")[:500] or None, credit=v["credit"], debit=v["debit"],
                            reference_number=v.get("reference_number"), utr=v.get("utr"),
                            cheque_number=v.get("cheque_number"), balance=v.get("balance"),
                            counterparty=v.get("counterparty"), recon_status="UNMATCHED", created_by=ctx.user.id,
                            **src)
        from app.services.bank import prepare_bank_txn
        prepare_bank_txn(t)
        t.txn_hash = row.row_hash  # includes in-file occurrence / template key
        db.add(t)
        db.flush()
        recompute(db, t)
        _hist(db, ctx.user, t, "IMPORT", None, remarks=f"Batch #{batch.id} row {row.source_row}")
        return t
    if fam == "FUEL":
        from app.services.operations import link_odometer, process_fuel
        t = FuelTransaction(provider_id=batch.provider_id, provider_transaction_id=v.get("provider_transaction_id"),
                            txn_datetime=v["txn_datetime"], txn_date=v["txn_date"],
                            vehicle_id=v.get("vehicle_number__id"), vehicle_number_raw=v.get("vehicle_number"),
                            fuel_station_id=v.get("station_name__id") or v.get("station_code__id"),
                            station_raw=v.get("station_name") or v.get("station_code"),
                            fuel_type_id=v.get("fuel_type__id"), fuel_type_raw=v.get("fuel_type"),
                            quantity=v["quantity"], rate=v.get("rate"), amount=v.get("amount"),
                            tax_amount=v.get("tax_amount"), discount_amount=v.get("discount_amount"),
                            total_amount=v.get("total_amount"), invoice_number=v.get("invoice_number"),
                            card_number=v.get("card_number"), odometer=v.get("odometer"), remarks=v.get("remarks"),
                            created_by=ctx.user.id, **src)
        if t.fuel_station_id is None and (v.get("station_code") or v.get("station_name")):
            from app.services.operations import match_station
            t.fuel_station_id = match_station(db, batch.provider_id, v.get("station_code"), v.get("station_name"))
        process_fuel(ctx, t, "IMPORT")
        t.txn_hash = row.row_hash
        db.add(t)
        db.flush()
        link_odometer(t)
        return t
    if fam == "TOLL":
        from app.services.operations import process_toll
        t = TollTransaction(provider_id=batch.provider_id, transaction_id=str(v["transaction_id"]),
                            txn_datetime=v["txn_datetime"], txn_date=v["txn_date"], txn_time=v.get("txn_time"),
                            vehicle_id=v.get("vehicle_number__id"), registration_raw=v.get("vehicle_number"),
                            fastag_id=v.get("fastag_id"), plaza_external_id_raw=v.get("plaza_id"),
                            plaza_code_raw=v.get("plaza_code"), plaza_name_raw=v.get("plaza_name"),
                            toll_plaza_id=v.get("plaza_name__id") or v.get("plaza_id__id") or v.get("plaza_code__id"),
                            amount=v["amount"], direction=v.get("direction"), lane=v.get("lane"),
                            txn_kind=v.get("txn_kind") or "DEBIT", remarks=v.get("remarks"), created_by=ctx.user.id,
                            **src)
        process_toll(ctx, t)
        db.add(t)
        db.flush()
        return t
    if fam == "MAINTENANCE":
        from app.services.vehicles import cost_center_for
        j = MaintenanceJobCard(job_card_number=v["job_card_number"], vehicle_id=v["vehicle_number__id"],
                               job_date=v["job_date"], maintenance_type_id=v["maintenance_type__id"],
                               odometer=v.get("odometer"), complaint=v.get("complaint"),
                               work_performed=v.get("work_performed"), vendor_id=v.get("vendor_name__id"),
                               parts_amount=v.get("parts_amount") or 0, labour_amount=v.get("labour_amount") or 0,
                               other_amount=v.get("other_amount") or 0, tax_amount=v.get("tax_amount") or 0,
                               invoice_number=v.get("invoice_number"), status="COMPLETED", created_by=ctx.user.id,
                               **src)
        j.cost_center_id = cost_center_for(db, j.vehicle_id, j.job_date)
        j.amount = j.parts_amount + j.labour_amount + j.other_amount
        j.total_amount = j.amount + j.tax_amount
        db.add(j)
        db.flush()
        return j
    if fam == "INSURANCE":
        from app.models.compliance import InsurancePolicy
        from app.services.renewals import compute_status
        p = InsurancePolicy(provider_id=batch.provider_id, vehicle_id=v.get("vehicle_number__id"),
                            policy_type=str(v["policy_type"]).upper(), policy_number=v["policy_number"],
                            start_date=v["start_date"], expiry_date=v["expiry_date"],
                            insured_amount=v.get("insured_amount"), premium=v.get("premium"),
                            tax_amount=v.get("tax_amount"), created_by=ctx.user.id)
        p.total_premium = (p.premium or 0) + (p.tax_amount or 0)
        p.renewal_status = compute_status(db, p.expiry_date)
        db.add(p)
        db.flush()
        return p
    if fam == "GPS":
        from app.services.vehicles import record_odometer
        r = record_odometer(ctx, v["vehicle_number__id"], v["odometer"], v["reading_datetime"], "GPS")
        return r
    raise BusinessError(f"Unsupported statement type {fam}")


def cancel_batch(ctx, batch_id: int) -> None:
    b = ctx.db.get(StatementImportBatch, batch_id)
    if not b or b.status not in ("PREVIEWED", "QUEUED"):
        raise BusinessError("Only previewed batches can be cancelled")
    b.status = "CANCELLED"
    audit(ctx.db, ctx.user, "CANCEL", "statement_import_batches", b.id)


def run_commit_job(batch_id: int, user_id: int) -> None:
    """Background execution for large files. Progress is committed in a separate session."""
    from app.core.security import load_user
    from app.database import SessionLocal, get_engine
    from app.services.masters import Ctx
    get_engine()
    db = SessionLocal()
    try:
        user = load_user(db, user_id)

        def progress(done, total):
            s2 = SessionLocal()
            try:
                b = s2.get(StatementImportBatch, batch_id)
                b.processed_rows = done
                s2.commit()
            finally:
                s2.close()

        commit_batch(Ctx(db, user), batch_id, progress_cb=progress)
        db.commit()
    except Exception as exc:  # noqa: BLE001
        db.rollback()
        log.exception("Import batch %s failed", batch_id)
        s2 = SessionLocal()
        try:
            b = s2.get(StatementImportBatch, batch_id)
            b.status, b.error_message = "FAILED", str(getattr(exc, "message", exc))[:2000]
            s2.commit()
        finally:
            s2.close()
    finally:
        db.close()
