"""Generic, metadata-driven CRUD engine.

Each screen is described once by a `MasterSpec` (see services/registry.py).
The same spec drives:

* REST endpoints  (/api/masters/{key} …) with server-side search, filters,
  sorting, pagination and CSV/XLSX export;
* validation/coercion of input (types, required, lengths, patterns,
  FK existence, configurable lookup values, uniqueness);
* audit trail (CREATE / UPDATE diff / ACTIVATE / DEACTIVATE / DELETE);
* child grids (e.g. invoice income allocations, job-card parts, template
  columns) saved atomically with their parent;
* row actions (renew, reclassify, resolve ...) implemented by domain services;
* the HTML form + list rendered by static/js/master.js in the house design.

Domain rules stay in domain services and are plugged in with hooks, so the
generic layer never contains business logic of its own.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any, Callable

from sqlalchemy import DateTime, String, and_, asc, cast, desc, func, or_, select
from sqlalchemy import inspect as sa_inspect
from sqlalchemy.orm import Session

from app.core.audit import audit, snapshot
from app.core.errors import BusinessError, Conflict, NotFound, PermissionDenied
from app.core.utils import (fmt_date, jsonable, now, parse_date, parse_datetime, parse_time, to_decimal)


# ═════════════════════════════ spec objects ═════════════════════════════
@dataclass
class F:
    name: str
    label: str
    type: str = "text"  # text textarea int decimal money bool date datetime time select lookup fk password json
    required: bool = False
    fk: str | None = None  # registry key of the referenced master
    lookup: str | None = None  # lookup_values.category
    choices: list[str] | None = None  # technical enums that carry program semantics
    section: str = "General"
    list: bool = False
    search: bool = False
    filter: bool = False
    upper: bool = False
    maxlen: int | None = None
    pattern: str | None = None
    pattern_msg: str | None = None
    readonly: bool = False
    readonly_on_edit: bool = False
    hint: str | None = None
    default: Any = None
    span: int = 1  # grid columns spanned in the form
    min: float | None = None
    max: float | None = None
    fk_filter: dict | None = None  # static option filter, e.g. {"provider_type": "BANK"}
    depends_on: str | None = None  # option filter from another field, "sub_category_id:cost_category_id"
    hidden: bool = False
    virtual: bool = False  # not a model column; passed to hooks only (e.g. password, role_ids)

    def meta(self) -> dict:
        d = {k: v for k, v in self.__dict__.items() if v not in (None, False) or k in ("required",)}
        d["name"] = self.name
        return d


@dataclass
class Child:
    key: str
    label: str
    model: Any
    fk_field: str
    fields: list[F]
    order_by: str = "id"
    min_rows: int = 0
    readonly: bool = False
    prepare: Callable | None = None  # (ctx, parent, child_obj) before insert, e.g. derive cost center


@dataclass
class Action:
    key: str
    label: str
    icon: str
    permission: str
    handler: Callable  # (ctx, obj, data) -> dict | None
    fields: list[F] = field(default_factory=list)
    confirm: str | None = None
    visible_when: dict | None = None  # {"status": ["OPEN"]}
    style: str = "edit"  # CSS flavour for the row button
    reason_required: bool = False


@dataclass
class MasterSpec:
    key: str
    model: Any
    title: str
    module: str
    icon: str = "ti-database"
    fields: list[F] = field(default_factory=list)
    children: list[Child] = field(default_factory=list)
    actions: list[Action] = field(default_factory=list)
    title_field: str = "name"
    code_field: str | None = "code"
    label_fields: tuple[str, ...] = ("code", "name")
    default_sort: str = "id"
    default_dir: str = "desc"
    date_field: str | None = None
    can_create: bool = True
    can_edit: bool = True
    can_delete: bool = True
    hard_delete: bool = False
    before_save: Callable | None = None  # (ctx, obj, data, is_new) -> None
    after_save: Callable | None = None  # (ctx, obj, data, is_new) -> None
    before_delete: Callable | None = None
    serialize_extra: Callable | None = None  # (ctx, obj, row_dict) -> None
    list_query: Callable | None = None  # (ctx, stmt, params) -> stmt
    branch_field: str | None = None
    branch_via_vehicle: str | None = None
    documents: bool = False
    entity_type: str | None = None
    group: str = "Masters"
    description: str | None = None
    fixed_filters: dict | None = None  # always applied (e.g. lookup category screens)
    status_field: str | None = None
    subtitle_fields: tuple[str, ...] = ()
    page_size: int = 25

    @property
    def etype(self) -> str:
        return self.entity_type or self.model.__tablename__

    @property
    def has_active(self) -> bool:
        return hasattr(self.model, "is_active")

    def field(self, name: str) -> F | None:
        return next((f for f in self.fields if f.name == name), None)

    def perm(self, action: str) -> str:
        return f"{self.module}.{action}"


@dataclass
class Ctx:
    db: Session
    user: Any
    spec: MasterSpec | None = None
    warnings: list[str] = field(default_factory=list)
    reason: str | None = None
    override: bool = False


REGISTRY: dict[str, MasterSpec] = {}


def register(spec: MasterSpec) -> MasterSpec:
    REGISTRY[spec.key] = spec
    return spec


def get_spec(key: str) -> MasterSpec:
    spec = REGISTRY.get(key)
    if not spec:
        raise NotFound(f"Screen '{key}'")
    return spec


# ═════════════════════════════ metadata ═════════════════════════════
def spec_meta(spec: MasterSpec, user) -> dict:
    return {
        "key": spec.key, "title": spec.title, "icon": spec.icon, "module": spec.module, "group": spec.group,
        "description": spec.description,
        "fields": [f.meta() for f in spec.fields],
        "children": [{"key": c.key, "label": c.label, "readonly": c.readonly, "min_rows": c.min_rows,
                      "fields": [f.meta() for f in c.fields]} for c in spec.children],
        "actions": [{"key": a.key, "label": a.label, "icon": a.icon, "confirm": a.confirm, "style": a.style,
                     "visible_when": a.visible_when, "reason_required": a.reason_required,
                     "fields": [f.meta() for f in a.fields], "allowed": user.has(a.permission)}
                    for a in spec.actions],
        "title_field": spec.title_field, "code_field": spec.code_field, "status_field": spec.status_field,
        "subtitle_fields": list(spec.subtitle_fields),
        "has_active": spec.has_active, "documents": spec.documents, "entity_type": spec.etype,
        "date_field": spec.date_field, "default_sort": spec.default_sort, "default_dir": spec.default_dir,
        "page_size": spec.page_size,
        "perms": {
            "view": user.has(spec.perm("view")),
            "create": spec.can_create and user.has(spec.perm("create")),
            "edit": spec.can_edit and user.has(spec.perm("edit")),
            "delete": spec.can_delete and user.has(spec.perm("delete")),
            "audit": user.has("audit.view"),
        },
    }


# ═════════════════════════════ coercion / validation ═════════════════════════════
def _coerce(db: Session, f: F, raw: Any) -> Any:
    if raw is None or (isinstance(raw, str) and raw.strip() == ""):
        return None
    t = f.type
    try:
        if t in ("text", "textarea", "password"):
            v = str(raw).strip()
            if f.upper:
                v = v.upper()
            if f.maxlen and len(v) > f.maxlen:
                raise BusinessError(f"{f.label}: maximum {f.maxlen} characters")
            if f.pattern and not re.fullmatch(f.pattern, v):
                raise BusinessError(f"{f.label}: {f.pattern_msg or 'invalid format'}")
            return v
        if t == "int":
            v = int(Decimal(str(raw).replace(",", "")))
        elif t in ("decimal", "money"):
            v = to_decimal(raw)
        elif t == "bool":
            return raw if isinstance(raw, bool) else str(raw).lower() in ("1", "true", "yes", "on", "y")
        elif t == "date":
            return parse_date(raw)
        elif t == "datetime":
            v = parse_datetime(raw)
            if v is None:
                raise ValueError(f"Invalid date/time '{raw}'")
            return v
        elif t == "time":
            return parse_time(raw)
        elif t == "fk":
            v = int(raw)
            target = get_spec(f.fk)
            if db.get(target.model, v) is None:
                raise BusinessError(f"{f.label}: selected record does not exist")
            return v
        elif t == "lookup":
            from app.models.org import LookupValue
            v = str(raw).strip().upper()
            ok = db.execute(select(LookupValue.id).where(LookupValue.category == f.lookup,
                                                         LookupValue.code == v)).first()
            if not ok:
                raise BusinessError(f"{f.label}: '{v}' is not a configured value")
            return v
        elif t == "select":
            v = str(raw).strip()
            if f.choices and v not in f.choices:
                raise BusinessError(f"{f.label}: invalid choice '{v}'")
            return v
        elif t == "json":
            return raw
        else:
            return raw
    except BusinessError:
        raise
    except (ValueError, TypeError, ArithmeticError) as exc:
        raise BusinessError(f"{f.label}: {exc}") from exc
    if f.min is not None and v is not None and v < f.min:
        raise BusinessError(f"{f.label}: must be ≥ {f.min}")
    if f.max is not None and v is not None and v > f.max:
        raise BusinessError(f"{f.label}: must be ≤ {f.max}")
    return v


def _apply_fields(db: Session, fields: list[F], obj: Any, data: dict, is_new: bool, errors: list) -> None:
    for f in fields:
        if f.readonly or f.virtual or (f.readonly_on_edit and not is_new) or f.type in ("password", "multi"):
            continue
        if f.name not in data:
            if is_new and f.default is not None and getattr(obj, f.name, None) is None:
                setattr(obj, f.name, f.default)
            if is_new and f.required and getattr(obj, f.name, None) in (None, ""):
                errors.append({"field": f.name, "message": f"{f.label} is required"})
            continue
        try:
            v = _coerce(db, f, data[f.name])
        except BusinessError as exc:
            errors.append({"field": f.name, "message": exc.message})
            continue
        if v is None and f.required:
            errors.append({"field": f.name, "message": f"{f.label} is required"})
            continue
        if v is None and f.type == "bool":
            v = False
        setattr(obj, f.name, v)


def _check_unique(db: Session, spec: MasterSpec, obj: Any, errors: list) -> None:
    """Friendly duplicate messages before the DB unique index fires."""
    table = spec.model.__table__
    groups = [[c.name] for c in table.columns if c.unique]
    for con in table.constraints:
        if con.__class__.__name__ == "UniqueConstraint":
            groups.append([c.name for c in con.columns])
    for cols in groups:
        vals = {c: getattr(obj, c, None) for c in cols}
        if any(v is None for v in vals.values()):
            continue
        q = select(spec.model.id).where(and_(*[getattr(spec.model, c) == v for c, v in vals.items()]))
        if obj.id:
            q = q.where(spec.model.id != obj.id)
        if db.execute(q).first():
            labels = ", ".join((spec.field(c).label if spec.field(c) else c) for c in cols)
            errors.append({"field": cols[-1], "message": f"Duplicate: another record already has this {labels}"})


# ═════════════════════════════ serialisation ═════════════════════════════
def _label_of(obj: Any, spec: MasterSpec) -> str:
    parts = [str(getattr(obj, f)) for f in spec.label_fields if getattr(obj, f, None) not in (None, "")]
    return " — ".join(parts) if parts else f"#{obj.id}"


def _resolve_labels(db: Session, fields: list[F], rows: list[dict]) -> None:
    from app.models.org import LookupValue
    for f in fields:
        if f.type == "fk" and f.fk:
            ids = {r[f.name] for r in rows if r.get(f.name)}
            if not ids:
                continue
            target = get_spec(f.fk)
            objs = db.execute(select(target.model).where(target.model.id.in_(ids))).scalars()
            labels = {o.id: _label_of(o, target) for o in objs}
            for r in rows:
                if r.get(f.name):
                    r[f"{f.name}__label"] = labels.get(r[f.name], f"#{r[f.name]}")
        elif f.type == "lookup":
            codes = {r[f.name] for r in rows if r.get(f.name)}
            if not codes:
                continue
            lbl = dict(db.execute(select(LookupValue.code, LookupValue.label).where(
                LookupValue.category == f.lookup, LookupValue.code.in_(codes))).all())
            for r in rows:
                if r.get(f.name):
                    r[f"{f.name}__label"] = lbl.get(r[f.name], r[f.name])


def serialize_many(ctx: Ctx, spec: MasterSpec, objs: list, with_children: bool = False) -> list[dict]:
    rows = []
    cols = [a.key for a in sa_inspect(spec.model).mapper.column_attrs]
    for o in objs:
        r = {c: jsonable(getattr(o, c)) for c in cols if "password" not in c}
        r["_label"] = _label_of(o, spec)
        rows.append(r)
    _resolve_labels(ctx.db, spec.fields, rows)
    if with_children:
        for o, r in zip(objs, rows):
            for ch in spec.children:
                kids = ctx.db.execute(select(ch.model).where(getattr(ch.model, ch.fk_field) == o.id)
                                      .order_by(getattr(ch.model, ch.order_by))).scalars().all()
                kr = [{a.key: jsonable(getattr(k, a.key)) for a in sa_inspect(ch.model).mapper.column_attrs}
                      for k in kids]
                _resolve_labels(ctx.db, ch.fields, kr)
                r[ch.key] = kr
    if spec.serialize_extra:
        for o, r in zip(objs, rows):
            spec.serialize_extra(ctx, o, r)
    return rows


# ═════════════════════════════ queries ═════════════════════════════
def _branch_filter(ctx: Ctx, spec: MasterSpec, stmt):
    bids = getattr(ctx.user, "branch_ids", None)
    if bids is None:
        return stmt
    if spec.branch_field:
        col = getattr(spec.model, spec.branch_field)
        return stmt.where(or_(col.is_(None), col.in_(bids)))
    if spec.branch_via_vehicle:
        from app.models.fleet import Vehicle
        col = getattr(spec.model, spec.branch_via_vehicle)
        return stmt.where(or_(col.is_(None), col.in_(select(Vehicle.id).where(Vehicle.branch_id.in_(bids)))))
    return stmt


def list_records(ctx: Ctx, spec: MasterSpec, params: dict, paginate: bool = True) -> dict:
    ctx.user.require(spec.perm("view"))
    M = spec.model
    stmt = select(M)
    for k, v in (spec.fixed_filters or {}).items():
        stmt = stmt.where(getattr(M, k) == v)
    active = params.get("active", "1")
    if spec.has_active and active in ("1", "0"):
        stmt = stmt.where(M.is_active.is_(active == "1"))
    q = (params.get("q") or "").strip()
    if q:
        exact = params.get("exact") in ("1", "true")
        conds = []
        for f in spec.fields:
            if f.search:
                col = getattr(M, f.name)
                conds.append(col == q if exact else cast(col, String).ilike(f"%{q}%"))
        if q.isdigit():
            conds.append(M.id == int(q))
        if conds:
            stmt = stmt.where(or_(*conds))
    for f in spec.fields:
        v = params.get(f"f_{f.name}")
        if v not in (None, ""):
            col = getattr(M, f.name)
            if f.type == "bool":
                stmt = stmt.where(col.is_(v in ("1", "true")))
            elif "," in str(v):
                stmt = stmt.where(col.in_(str(v).split(",")))
            else:
                stmt = stmt.where(col == (int(v) if f.type in ("fk", "int") else v))
    if spec.date_field:
        col = getattr(M, spec.date_field)
        if params.get("date_from"):
            stmt = stmt.where(col >= parse_date(params["date_from"]))
        if params.get("date_to"):
            d = parse_date(params["date_to"])
            if isinstance(col.type, DateTime):
                stmt = stmt.where(col < datetime.combine(d + timedelta(days=1), datetime.min.time()))
            else:
                stmt = stmt.where(col <= d)
    stmt = _branch_filter(ctx, spec, stmt)
    if spec.list_query:
        stmt = spec.list_query(ctx, stmt, params)
    total = ctx.db.execute(select(func.count()).select_from(stmt.order_by(None).subquery())).scalar_one()
    sort = params.get("sort") or spec.default_sort
    if not hasattr(M, sort):
        sort = spec.default_sort
    direction = params.get("dir") or spec.default_dir
    order = desc(getattr(M, sort)) if direction == "desc" else asc(getattr(M, sort))
    stmt = stmt.order_by(order, desc(M.id))
    page = max(int(params.get("page") or 1), 1)
    size = min(max(int(params.get("size") or spec.page_size), 1), 500)
    if paginate:
        stmt = stmt.offset((page - 1) * size).limit(size)
    objs = ctx.db.execute(stmt).scalars().all()
    return {"items": serialize_many(ctx, spec, objs), "total": total, "page": page, "size": size}


def get_obj(ctx: Ctx, spec: MasterSpec, rid: int):
    obj = ctx.db.get(spec.model, rid)
    if obj is None:
        raise NotFound(spec.title)
    bids = getattr(ctx.user, "branch_ids", None)
    if bids is not None:
        if spec.branch_field and getattr(obj, spec.branch_field) not in (None, *bids):
            raise PermissionDenied("Record belongs to a branch you cannot access")
    return obj


def get_record(ctx: Ctx, spec: MasterSpec, rid: int) -> dict:
    ctx.user.require(spec.perm("view"))
    return serialize_many(ctx, spec, [get_obj(ctx, spec, rid)], with_children=True)[0]


def options(ctx: Ctx, spec: MasterSpec, q: str = "", limit: int = 50, filters: dict | None = None,
            include_id: int | None = None) -> list[dict]:
    M = spec.model
    stmt = select(M)
    for k, v in (spec.fixed_filters or {}).items():
        stmt = stmt.where(getattr(M, k) == v)
    if spec.has_active:
        stmt = stmt.where(M.is_active.is_(True))
    for k, v in (filters or {}).items():
        if hasattr(M, k) and v not in (None, ""):
            stmt = stmt.where(getattr(M, k).in_(str(v).split(",")) if "," in str(v) else getattr(M, k) == v)
    if q:
        conds = [cast(getattr(M, f), String).ilike(f"%{q}%") for f in spec.label_fields if hasattr(M, f)]
        stmt = stmt.where(or_(*conds))
    stmt = _branch_filter(ctx, spec, stmt)
    first = spec.label_fields[0] if spec.label_fields and hasattr(M, spec.label_fields[0]) else "id"
    objs = list(ctx.db.execute(stmt.order_by(getattr(M, first)).limit(limit)).scalars())
    if include_id and not any(o.id == include_id for o in objs):
        extra = ctx.db.get(M, include_id)
        if extra:
            objs.append(extra)
    return [{"id": o.id, "label": _label_of(o, spec)} for o in objs]


def _save_children(ctx: Ctx, spec: MasterSpec, obj: Any, data: dict, errors: list) -> dict:
    out = {}
    for ch in spec.children:
        if ch.readonly or ch.key not in data:
            continue
        rows = data.get(ch.key) or []
        if len(rows) < ch.min_rows:
            errors.append({"field": ch.key, "message": f"At least {ch.min_rows} {ch.label} line(s) required"})
            continue
        existing = ctx.db.execute(select(ch.model).where(getattr(ch.model, ch.fk_field) == obj.id)).scalars().all()
        out[ch.key] = {"old": [snapshot(e) for e in existing]}
        for e in existing:
            ctx.db.delete(e)
        ctx.db.flush()
        new_objs = []
        for i, r in enumerate(rows, 1):
            kid = ch.model()
            setattr(kid, ch.fk_field, obj.id)
            row_err: list = []
            _apply_fields(ctx.db, ch.fields, kid, r, True, row_err)
            for e in row_err:
                errors.append({"field": f"{ch.key}[{i}].{e['field']}", "message": f"{ch.label} line {i}: {e['message']}"})
            if row_err:
                continue
            if ch.prepare:
                ch.prepare(ctx, obj, kid)
            ctx.db.add(kid)
            new_objs.append(kid)
        out[ch.key]["new"] = new_objs
    return out


def save_record(ctx: Ctx, spec: MasterSpec, data: dict, rid: int | None = None) -> dict:
    is_new = rid is None
    ctx.spec = spec
    if is_new:
        if not spec.can_create:
            raise PermissionDenied("Records cannot be created on this screen")
        ctx.user.require(spec.perm("create"))
        obj = spec.model()
        old = None
    else:
        if not spec.can_edit:
            raise PermissionDenied("Records cannot be edited on this screen")
        ctx.user.require(spec.perm("edit"))
        obj = get_obj(ctx, spec, rid)
        old = snapshot(obj)
    errors: list = []
    _apply_fields(ctx.db, spec.fields, obj, data, is_new, errors)
    if errors:
        raise BusinessError("Please correct the highlighted fields", "VALIDATION", errors, 422)
    if spec.before_save:
        spec.before_save(ctx, obj, data, is_new)
    _check_unique(ctx.db, spec, obj, errors)
    if errors:
        raise Conflict("Duplicate value", "DUPLICATE", errors)
    if hasattr(obj, "updated_by"):
        obj.updated_by = ctx.user.id
        if is_new:
            obj.created_by = ctx.user.id
    if is_new:
        ctx.db.add(obj)
    ctx.db.flush()
    kids = _save_children(ctx, spec, obj, data, errors)
    if errors:
        raise BusinessError("Please correct the highlighted lines", "VALIDATION", errors, 422)
    ctx.db.flush()
    if spec.after_save:
        spec.after_save(ctx, obj, data, is_new)
    ctx.db.flush()
    new = snapshot(obj)
    child_new = {k: [snapshot(x) for x in v["new"]] for k, v in kids.items()}
    if child_new:
        new = {**new, **{f"__{k}": v for k, v in child_new.items()}}
        if old is not None:
            old = {**old, **{f"__{k}": v["old"] for k, v in kids.items()}}
    audit(ctx.db, ctx.user, "CREATE" if is_new else "UPDATE", spec.etype, obj.id, old, new, reason=ctx.reason)
    rec = serialize_many(ctx, spec, [obj], with_children=True)[0]
    rec["_warnings"] = ctx.warnings
    return rec


def set_active(ctx: Ctx, spec: MasterSpec, rid: int, active: bool, reason: str | None = None) -> dict:
    ctx.user.require(spec.perm("edit"))
    if not spec.has_active:
        raise BusinessError("This record type has no active flag")
    obj = get_obj(ctx, spec, rid)
    obj.is_active = active
    obj.deleted_at = None if active else obj.deleted_at
    obj.updated_by = ctx.user.id
    audit(ctx.db, ctx.user, "ACTIVATE" if active else "DEACTIVATE", spec.etype, rid, reason=reason)
    ctx.db.flush()
    return serialize_many(ctx, spec, [obj])[0]


def delete_record(ctx: Ctx, spec: MasterSpec, rid: int, reason: str | None = None) -> None:
    ctx.user.require(spec.perm("delete"))
    if not spec.can_delete:
        raise PermissionDenied("Records on this screen cannot be deleted")
    obj = get_obj(ctx, spec, rid)
    if spec.before_delete:
        spec.before_delete(ctx, obj)
    old = snapshot(obj)
    if spec.has_active and not spec.hard_delete:
        obj.is_active = False
        obj.deleted_at = now()
        obj.updated_by = ctx.user.id
        audit(ctx.db, ctx.user, "SOFT_DELETE", spec.etype, rid, old=old, reason=reason)
    else:
        audit(ctx.db, ctx.user, "DELETE", spec.etype, rid, old=old, reason=reason)
        ctx.db.delete(obj)
    ctx.db.flush()


def run_action(ctx: Ctx, spec: MasterSpec, rid: int, action_key: str, data: dict) -> dict:
    act = next((a for a in spec.actions if a.key == action_key), None)
    if not act:
        raise NotFound("Action")
    ctx.user.require(act.permission)
    obj = get_obj(ctx, spec, rid)
    errors: list = []
    clean: dict = {}
    for f in act.fields:
        try:
            v = _coerce(ctx.db, f, data.get(f.name))
        except BusinessError as exc:
            errors.append({"field": f.name, "message": exc.message})
            continue
        if v is None and f.required:
            errors.append({"field": f.name, "message": f"{f.label} is required"})
        clean[f.name] = v
    if act.reason_required and not (data.get("reason") or clean.get("reason")):
        errors.append({"field": "reason", "message": "Reason is required"})
    if errors:
        raise BusinessError("Please correct the highlighted fields", "VALIDATION", errors, 422)
    clean.setdefault("reason", data.get("reason"))
    ctx.reason = clean.get("reason")
    ctx.override = bool(data.get("override"))
    result = act.handler(ctx, obj, clean) or {}
    ctx.db.flush()
    result.setdefault("record", serialize_many(ctx, spec, [obj], with_children=True)[0])
    result.setdefault("warnings", ctx.warnings)
    return result


# ═════════════════════════════ export ═════════════════════════════
def export_rows(ctx: Ctx, spec: MasterSpec, params: dict) -> tuple[list[str], list[list]]:
    data = list_records(ctx, spec, {**params, "size": 500}, paginate=False)
    cols = [f for f in spec.fields if (f.list or f.search) and not f.virtual and f.type not in ("password", "multi")]
    headers = [f.label for f in cols]
    rows = []
    for r in data["items"]:
        row = []
        for f in cols:
            v = r.get(f"{f.name}__label", r.get(f.name))
            if f.type in ("date", "datetime") and v:
                v = fmt_date(parse_datetime(v) if f.type == "datetime" else parse_date(v))
            elif f.type == "bool":
                v = "Yes" if v else "No"
            row.append(v)
        rows.append(row)
    return headers, rows


def to_csv(headers: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerows(rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(title: str, headers: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", title)[:31] or "Report"
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A5F")
    for r in rows:
        ws.append([float(x) if isinstance(x, Decimal) else x for x in r])
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col[:200])
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(title: str, headers: list[str], rows: list[list], subtitle: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    ss = getSampleStyleSheet()
    cell = ss["BodyText"].clone("cell", fontSize=6.5, leading=8)
    data = [[Paragraph(f"<b>{h}</b>", cell) for h in headers]]
    for r in rows[:5000]:
        data.append([Paragraph("" if v is None else str(v).replace("&", "&amp;").replace("<", "&lt;"), cell)
                     for v in r])
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story = [Paragraph(title, ss["Title"]), Paragraph(subtitle or f"Generated {fmt_date(now())}", ss["Normal"]),
             Spacer(1, 8), t]
    doc.build(story)
    return buf.getvalue()


def today_str() -> str:
    return fmt_date(date.today())
